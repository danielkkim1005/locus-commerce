"""Turn out/round2_sheet.csv into one .xlsx per labeler, with dropdowns.

Round 1 collected every label as free text and paid for it: `is_commerce` came
back as yes/y/true/1, `category` came back with several categories in one cell,
and merge_labels.py needed a normalizer per field plus a fix for labeler identity.
All of that is a data entry problem, not a labeling problem.

Dropdowns remove it at the source. Every column is a closed list except `notes`,
which stays free text because that is where the reasoning goes.

CSV cannot carry data validation, so these are .xlsx. Google Sheets preserves
validation on import: upload the file, let Drive convert it, and the dropdowns
are still there.

    pip install openpyxl
    python make_sheets.py

Writes out/round2_<name>.xlsx, one per --labelers name, each with:

    a Labels tab      the 75 rows, dropdowns on every coded column
    a Guide tab       the tiers, the decision procedure, the reason codes
    labeler prefilled on every row, so identity cannot drift mid-sheet

Options:

    --sheet      the master CSV (default: out/round2_sheet.csv)
    --out        output folder (default: out)
    --labelers   comma separated (default: ainesh,daniel,sachin)
"""

import argparse
import csv
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Must match definition_v2.md. If you edit one, edit the other.
TIERS = ["core", "adjacent", "not"]
ADJACENT_REASONS = ["gov-market", "revenue", "land-use", "general-law",
                    "counterparty", "procedure", "definition", "sometimes-business"]
CATEGORIES = ["entry", "operation", "exit", "other"]
UNRESOLVABLE = ["truncated", "no-context"]
STEPS = ["1", "2", "3", "4", "5"]

# column -> (allowed values, prompt shown when the cell is selected)
VALIDATED = {
    "tier": (TIERS,
             "core = duty/prohibition/sanction/eligibility on a commercial party. "
             "adjacent = commercial subject, but pick a reason. "
             "not = fails step 2 or 3. Leave blank if unresolvable."),
    "adjacent_reason": (ADJACENT_REASONS,
                        "Required when tier = adjacent. Leave blank otherwise. "
                        "If nothing fits, say so in notes - that is a finding."),
    "category": (CATEGORIES,
                 "entry = getting started. operation = while running. "
                 "exit = losing the ability to operate. other = commercial but none of those."),
    "unresolvable": (UNRESOLVABLE,
                     "truncated = text cut off mid-provision. "
                     "no-context = the chunk never names what it governs. "
                     "Set this and leave tier blank."),
    "step": (STEPS, "Which step of the decision procedure decided it (1-5)."),
}
FREE_TEXT = ["notes"]
READ_ONLY = ["row_id", "chunk_id", "header", "text", "text_truncated"]

GUIDE = [
    ("The three tiers", None),
    ("core", "The provision imposes a duty, prohibition, sanction, or eligibility "
             "condition on a party because that party engages in commerce."),
    ("adjacent", "Commercial subject matter, but one of the reason codes below means "
                 "it is not a conduct rule on a commercial party. Not a bucket for "
                 "hard rows - pick a reason or it is not adjacent."),
    ("not", "Fails step 2 or step 3. No commercial subject, or generally applicable "
            "law that happens to bind businesses too."),
    ("", None),
    ("Decision procedure - stop at the first step that answers", None),
    ("1", "Can you tell what is regulated? If the text is cut off, or never names the "
          "activity or party it governs, set unresolvable and leave tier blank. "
          "Do not guess from the header."),
    ("2", "Is any commercial subject named - business, trade, occupation, licensee, "
          "vendor, franchisee, commercial premises, commercial use of land? If no: not."),
    ("3", "Would a non-commercial actor doing the same thing be bound? If yes it is "
          "generally applicable: not. Unless the provision is written only for "
          "licensees or only for commercial premises - then continue to step 5."),
    ("4", "Is the regulated party the business itself? If the duty falls on someone "
          "dealing with a business - customer, employee, counterparty - then "
          "adjacent / counterparty, not core."),
    ("5", "Does it impose a duty, prohibition, sanction, or eligibility condition on "
          "that commercial party? Yes: core. Otherwise: adjacent, and pick a reason."),
    ("", None),
    ("Adjacent reason codes", None),
    ("gov-market", "Government as buyer, seller, or grantor rather than regulator. "
                   "Procurement, competitive bidding, franchise grants and reserved "
                   "rights, municipal utility tariffs, prices the city sets."),
    ("revenue", "Raises money rather than governing conduct. Occupation and license "
                "taxes, exemptions, fee schedules. BUT: if payment is a condition of "
                "holding the license, that is core (step 5)."),
    ("land-use", "Effect on commerce runs through permitted use of land, not a rule on "
                 "a trade. Permitted-use tables, commercial districts, setbacks. BUT: "
                 "home occupation or short-term rental conditions on the operator are core."),
    ("general-law", "Generally applicable law that reached step 5 only because it was "
                    "scoped to licensees. All-signs sign codes, a licensing division "
                    "enforcing ordinary traffic law. BUT: commercial signage, "
                    "advertising, false advertising are core."),
    ("counterparty", "The duty falls on someone dealing with a business rather than on "
                     "the business. Customer obligations, employee conduct provisions "
                     "written against the employee."),
    ("procedure", "Administrative machinery of an identifiably commercial chapter. How "
                  "to apply, what the clerk does, hearings, renewal mechanics, "
                  "inspection steps. If the chapter's subject is not identifiable, "
                  "that is step 1, not this."),
    ("definition", "A definitions section for an identifiably commercial chapter."),
    ("sometimes-business", "The activity is commercial only sometimes and the text does "
                           "not say which. Kennels, day care, massage, towing. If the "
                           "text specifies a commercial operator it is core. Provisional "
                           "code - only one round-1 row hit it."),
    ("", None),
    ("House rules", None),
    ("", "Judge the text in front of you against definition_v2.md. Do not look up the "
         "jurisdiction or the rest of the code, and do not check what anyone else said."),
    ("", "notes is the most valuable column. Write why whenever the call was hard. "
         "These become the edge-case log, and they are what edits the definition."),
    ("", "text_truncated = yes means OUR export cut the text at 6000 characters, not "
         "LOCUS. Only 1 of 75 rows. Do not call that unresolvable/truncated."),
    ("", "Do not edit row_id, chunk_id, header, or text. chunk_id is what the merge "
         "joins on, and an edited id silently drops the row."),
]

HEAD_FILL = PatternFill("solid", fgColor="DDE5EE")
FILLIN_FILL = PatternFill("solid", fgColor="FFF6DC")
RO_FONT = Font(color="808080")


def write_guide(ws) -> None:
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 108
    row = 1
    for key, text in GUIDE:
        a = ws.cell(row=row, column=1, value=key)
        if text is None:                      # a section heading
            a.font = Font(bold=True, size=12)
        else:
            a.font = Font(bold=True)
            b = ws.cell(row=row, column=2, value=text)
            b.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    ws.freeze_panes = "A2"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", default="out/round2_sheet.csv")
    ap.add_argument("--out", default="out")
    ap.add_argument("--labelers", default="ainesh,daniel,sachin")
    args = ap.parse_args()

    src = pathlib.Path(args.sheet)
    if not src.exists():
        raise SystemExit(f"{src} not found - run export_round2.py --hf first")
    with src.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    cols = list(rows[0].keys())
    if "labeler" not in cols:
        raise SystemExit(f"{src} has no labeler column")

    out = pathlib.Path(args.out)
    names = [n.strip() for n in args.labelers.split(",") if n.strip()]

    for name in names:
        wb = Workbook()
        ws = wb.active
        ws.title = "Labels"

        for j, col in enumerate(cols, start=1):
            c = ws.cell(row=1, column=j, value=col)
            c.font = Font(bold=True)
            c.fill = HEAD_FILL
            c.alignment = Alignment(vertical="center")

        for i, row in enumerate(rows, start=2):
            for j, col in enumerate(cols, start=1):
                # Prefill labeler on every row. Round 1 needed a fix in
                # merge_labels.py for labeler identity; a prefilled column cannot
                # drift, be left blank, or be spelled three ways in one sheet.
                value = name if col == "labeler" else row[col]
                c = ws.cell(row=i, column=j, value=value)
                if col in READ_ONLY:
                    c.font = RO_FONT
                if col in VALIDATED or col in FREE_TEXT:
                    c.fill = FILLIN_FILL
                if col == "text":
                    c.alignment = Alignment(wrap_text=True, vertical="top")

        last = len(rows) + 1
        for col, (allowed, prompt) in VALIDATED.items():
            letter = get_column_letter(cols.index(col) + 1)
            # The list goes inline in the formula rather than in a lookup range, so
            # there is no helper tab for someone to delete, and Sheets keeps it on
            # import. Excel caps the inline list at 255 characters; the longest here
            # is adjacent_reason at about 100.
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(allowed) + '"',
                allow_blank=True,
                showDropDown=False,   # False means DO show it; this flag is inverted
                showErrorMessage=True,
                errorTitle="Not an allowed value",
                error="Pick from the dropdown. If nothing fits, leave it blank and "
                      "explain in notes.",
                promptTitle=col,
                prompt=prompt,
                showInputMessage=True,
            )
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{last}")

        widths = {"row_id": 6, "chunk_id": 12, "header": 46, "text": 90,
                  "text_truncated": 9, "tier": 11, "adjacent_reason": 18,
                  "category": 11, "unresolvable": 13, "step": 6, "notes": 46,
                  "labeler": 10}
        for j, col in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(j)].width = widths.get(col, 14)
        ws.freeze_panes = "C2"          # keep row_id and chunk_id visible while scrolling
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

        write_guide(wb.create_sheet("Guide"))

        path = out / f"round2_{name}.xlsx"
        wb.save(path)
        print(f"wrote {path}  ({len(rows)} rows, labeler={name})")

    print(f"\n{len(names)} sheets, identical rows, same order.")
    print("Upload each to Drive and let it convert; the dropdowns survive the import.")
    print("Share each copy with that one person only. Keep out/round2_key.csv to yourself.")


if __name__ == "__main__":
    main()
